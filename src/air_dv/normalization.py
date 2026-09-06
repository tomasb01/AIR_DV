"""Normalization of source formats into an extractor-independent document model."""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from dataclasses import replace
from pathlib import Path
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from air_dv.models import Block, BlockType, DocumentSummary, NormalizedDocument, SourceLocation


_HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.+?)(?:\s+#+)?\s*$")
_LIST_ITEM_PATTERN = re.compile(r"^\s*(?:[-*+]\s+|\d+[.)]\s+)(.+)$")
_IMAGE_PATTERN = re.compile(r"!\[[^]]*\]\([^)]*\)")
_OBJECT_PLACEHOLDER_PATTERN = re.compile(
    r"<!--\s*(?:image|picture|object|drawing)[^>]*-->", re.IGNORECASE
)
_WORD_MARKDOWN_FORMATTING_PATTERN = re.compile(r"[*_`#]")
_WHITESPACE_PATTERN = re.compile(r"\s+")


class WordSourceLocator:
    """Map extracted Word blocks to stable paragraph positions in the source DOCX."""

    _namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

    def attach_locations(self, path: Path, blocks: tuple[Block, ...]) -> tuple[Block, ...]:
        """Attach Word paragraph positions when a normalized block matches source text exactly."""

        paragraphs = self._read_paragraphs(path)
        if not paragraphs:
            return blocks

        next_paragraph = 0
        located_blocks: list[Block] = []
        for block in blocks:
            source_text = self._canonical_text(block.text)
            match_index = self._find_match(paragraphs, source_text, next_paragraph)
            if match_index is None:
                located_blocks.append(block)
                continue
            paragraph_index, _ = paragraphs[match_index]
            located_blocks.append(
                replace(
                    block,
                    location=replace(
                        block.location,
                        label="Word document",
                        line_start=None,
                        line_end=None,
                        paragraph_index=paragraph_index,
                    ),
                )
            )
            next_paragraph = match_index + 1
        return tuple(located_blocks)

    @classmethod
    def _read_paragraphs(cls, path: Path) -> list[tuple[int, str]]:
        try:
            with ZipFile(path) as archive:
                document = ElementTree.fromstring(archive.read("word/document.xml"))
        except (BadZipFile, ElementTree.ParseError, KeyError, OSError):
            return []

        paragraphs = []
        for paragraph_index, paragraph in enumerate(document.findall(".//w:body/w:p", cls._namespace), start=1):
            text = "".join(node.text or "" for node in paragraph.findall(".//w:t", cls._namespace))
            canonical_text = cls._canonical_text(text)
            if canonical_text:
                paragraphs.append((paragraph_index, canonical_text))
        return paragraphs

    @staticmethod
    def _canonical_text(text: str) -> str:
        without_formatting = _WORD_MARKDOWN_FORMATTING_PATTERN.sub("", text)
        return _WHITESPACE_PATTERN.sub(" ", without_formatting).strip().casefold()

    @staticmethod
    def _find_match(
        paragraphs: list[tuple[int, str]], source_text: str, start_index: int
    ) -> int | None:
        if not source_text:
            return None
        for index in range(start_index, len(paragraphs)):
            _, paragraph_text = paragraphs[index]
            if source_text == paragraph_text:
                return index
        return None


class PdfSourceLocator:
    """Map extracted PDF blocks to source pages when extractable page text is available."""

    def attach_locations(self, path: Path, blocks: tuple[Block, ...]) -> tuple[Block, ...]:
        """Attach the best available source page to each normalized block."""

        page_texts = self._read_page_texts(path)
        if not page_texts:
            return blocks

        current_page = 1
        located_blocks: list[Block] = []
        for block in blocks:
            source_text = WordSourceLocator._canonical_text(block.text)
            matched_page = self._find_page(page_texts, source_text, current_page)
            if matched_page is not None:
                current_page = matched_page
            located_blocks.append(
                replace(
                    block,
                    location=replace(
                        block.location,
                        label="PDF document",
                        line_start=None,
                        line_end=None,
                        page_number=current_page,
                    ),
                )
            )
        return tuple(located_blocks)

    @staticmethod
    def _read_page_texts(path: Path) -> list[str]:
        try:
            reader = PdfReader(path)
            return [WordSourceLocator._canonical_text(page.extract_text() or "") for page in reader.pages]
        except Exception:
            return []

    @staticmethod
    def _find_page(page_texts: list[str], source_text: str, start_page: int) -> int | None:
        if not source_text:
            return None
        for page_number in range(start_page, len(page_texts) + 1):
            if source_text in page_texts[page_number - 1]:
                return page_number
        return None


class MarkdownNormalizer:
    """Read Markdown while preserving the semantic structures needed by AIR-DV checks."""

    file_type = "md"

    def normalize_file(self, path: str | Path) -> NormalizedDocument:
        """Normalize a UTF-8 Markdown file from disk."""

        source_path = Path(path)
        return self.normalize_text(source_path.name, source_path.read_text(encoding="utf-8"))

    def normalize_text(
        self,
        filename: str,
        content: str,
        *,
        source_label: str = "Source Markdown",
    ) -> NormalizedDocument:
        """Normalize Markdown text and retain source line locations for each block."""

        blocks: list[Block] = []
        lines = content.splitlines()
        index = 0

        while index < len(lines):
            line = lines[index]
            line_number = index + 1

            if not line.strip():
                index += 1
                continue

            heading_match = _HEADING_PATTERN.match(line)
            if heading_match:
                blocks.append(
                    Block(
                        type=BlockType.HEADING,
                        text=heading_match.group(2),
                        heading_level=len(heading_match.group(1)),
                        location=SourceLocation(label=source_label, line_start=line_number),
                    )
                )
                index += 1
                continue

            if line.lstrip().startswith("|"):
                index = self._collect_table(lines, index, blocks, source_label)
                continue

            object_match = _OBJECT_PLACEHOLDER_PATTERN.search(line)
            if object_match:
                blocks.append(
                    Block(
                        type=BlockType.UNSUPPORTED_OBJECT,
                        text=object_match.group(0),
                        location=SourceLocation(label=source_label, line_start=line_number),
                    )
                )
                index += 1
                continue

            image_match = _IMAGE_PATTERN.search(line)
            if image_match:
                blocks.append(
                    Block(
                        type=BlockType.IMAGE,
                        text=image_match.group(0),
                        location=SourceLocation(label=source_label, line_start=line_number),
                    )
                )
                index += 1
                continue

            list_match = _LIST_ITEM_PATTERN.match(line)
            if list_match:
                blocks.append(
                    Block(
                        type=BlockType.LIST_ITEM,
                        text=list_match.group(1),
                        location=SourceLocation(label=source_label, line_start=line_number),
                    )
                )
                index += 1
                continue

            index = self._collect_paragraph(lines, index, blocks, source_label)

        summary = DocumentSummary(
            filename=filename,
            file_type=self.file_type,
            extraction_succeeded=True,
        )
        return NormalizedDocument(document=summary, content=content, blocks=tuple(blocks))

    @staticmethod
    def _collect_table(
        lines: list[str], index: int, blocks: list[Block], source_label: str
    ) -> int:
        start_index = index
        table_lines: list[str] = []
        while index < len(lines) and lines[index].lstrip().startswith("|"):
            table_lines.append(lines[index])
            index += 1

        blocks.append(
            Block(
                type=BlockType.TABLE,
                text="\n".join(table_lines),
                location=SourceLocation(
                    label=source_label,
                    line_start=start_index + 1,
                    line_end=index,
                ),
            )
        )
        return index

    @staticmethod
    def _collect_paragraph(
        lines: list[str], index: int, blocks: list[Block], source_label: str
    ) -> int:
        start_index = index
        paragraph_lines: list[str] = []

        while index < len(lines):
            line = lines[index]
            if not line.strip() or _HEADING_PATTERN.match(line) or line.lstrip().startswith("|"):
                break
            if (
                _LIST_ITEM_PATTERN.match(line)
                or _IMAGE_PATTERN.search(line)
                or _OBJECT_PLACEHOLDER_PATTERN.search(line)
            ):
                break
            paragraph_lines.append(line)
            index += 1

        if paragraph_lines:
            blocks.append(
                Block(
                    type=BlockType.PARAGRAPH,
                    text="\n".join(paragraph_lines),
                    location=SourceLocation(
                        label=source_label,
                        line_start=start_index + 1,
                        line_end=index,
                    ),
                )
            )
            return index

        return index + 1


class DoclingWordNormalizer:
    """Normalize Word documents through the locally installed Docling CLI."""

    file_type = "docx"

    def __init__(self, executable: str = "docling") -> None:
        self.executable = executable
        self._markdown_normalizer = MarkdownNormalizer()

    def normalize_file(self, path: str | Path) -> NormalizedDocument:
        """Extract a Word file to Markdown and retain its original identity."""

        source_path = Path(path)
        if source_path.suffix.lower() != ".docx":
            raise ValueError("DoclingWordNormalizer only accepts .docx files")

        if shutil.which(self.executable) is None:
            return self._failed_document(
                source_path,
                f"Required extractor '{self.executable}' is not available on PATH.",
            )

        with tempfile.TemporaryDirectory(prefix="air-dv-docling-") as output_dir:
            command = [
                self.executable,
                "convert",
                str(source_path),
                "--to",
                "md",
                "--output",
                output_dir,
                "--image-export-mode",
                "placeholder",
            ]
            completed = subprocess.run(command, capture_output=True, text=True, check=False)
            output_path = Path(output_dir, f"{source_path.stem}.md")

            if completed.returncode != 0:
                detail = completed.stderr.strip() or completed.stdout.strip() or "Unknown extractor error."
                return self._failed_document(source_path, f"Docling extraction failed: {detail}")
            if not output_path.is_file():
                return self._failed_document(
                    source_path,
                    "Docling completed without producing a Markdown output file.",
                )

            extracted_markdown = output_path.read_text(encoding="utf-8")

        normalized_markdown = self._markdown_normalizer.normalize_text(
            source_path.name,
            extracted_markdown,
            source_label="Normalized Word content",
        )
        blocks = WordSourceLocator().attach_locations(source_path, normalized_markdown.blocks)
        return NormalizedDocument(
            document=DocumentSummary(
                filename=source_path.name,
                file_type=self.file_type,
                extraction_succeeded=True,
                extraction_notes=("Extracted with Docling.",),
            ),
            content=normalized_markdown.content,
            blocks=blocks,
        )

    def _failed_document(self, source_path: Path, note: str) -> NormalizedDocument:
        return NormalizedDocument(
            document=DocumentSummary(
                filename=source_path.name,
                file_type=self.file_type,
                extraction_succeeded=False,
                extraction_notes=(note,),
            ),
            content="",
            blocks=(),
        )


class DoclingPdfNormalizer:
    """Normalize PDF documents through Docling without embedding image payloads."""

    file_type = "pdf"

    def __init__(self, executable: str = "docling") -> None:
        self.executable = executable
        self._markdown_normalizer = MarkdownNormalizer()

    def normalize_file(self, path: str | Path) -> NormalizedDocument:
        """Extract a PDF to Markdown and retain source-page context where possible."""

        source_path = Path(path)
        if source_path.suffix.lower() != ".pdf":
            raise ValueError("DoclingPdfNormalizer only accepts .pdf files")
        if shutil.which(self.executable) is None:
            return self._failed_document(
                source_path, f"Required extractor '{self.executable}' is not available on PATH."
            )

        with tempfile.TemporaryDirectory(prefix="air-dv-docling-") as output_dir:
            command = [
                self.executable,
                "convert",
                str(source_path),
                "--to",
                "md",
                "--output",
                output_dir,
                "--image-export-mode",
                "placeholder",
            ]
            completed = subprocess.run(command, capture_output=True, text=True, check=False)
            output_path = Path(output_dir, f"{source_path.stem}.md")
            if completed.returncode != 0:
                detail = completed.stderr.strip() or completed.stdout.strip() or "Unknown extractor error."
                return self._failed_document(source_path, f"Docling extraction failed: {detail}")
            if not output_path.is_file():
                return self._failed_document(
                    source_path, "Docling completed without producing a Markdown output file."
                )
            extracted_markdown = output_path.read_text(encoding="utf-8")

        normalized_markdown = self._markdown_normalizer.normalize_text(
            source_path.name, extracted_markdown, source_label="Normalized PDF content"
        )
        ocr_warning_count = (completed.stdout + completed.stderr).count("RapidOCR returned empty result")
        notes = ["Extracted with Docling using image placeholders."]
        if ocr_warning_count:
            notes.append(f"Docling reported {ocr_warning_count} OCR warning(s).")
        return NormalizedDocument(
            document=DocumentSummary(
                filename=source_path.name,
                file_type=self.file_type,
                extraction_succeeded=True,
                extraction_notes=tuple(notes),
                source_metadata=(
                    ("page_count", str(self._page_count(source_path))),
                    ("visual_placeholder_count", str(extracted_markdown.count("<!--"))),
                    ("ocr_warning_count", str(ocr_warning_count)),
                ),
            ),
            content=normalized_markdown.content,
            blocks=PdfSourceLocator().attach_locations(source_path, normalized_markdown.blocks),
        )

    @staticmethod
    def _page_count(path: Path) -> int:
        try:
            return len(PdfReader(path).pages)
        except Exception:
            return 0

    def _failed_document(self, source_path: Path, note: str) -> NormalizedDocument:
        return NormalizedDocument(
            document=DocumentSummary(
                filename=source_path.name,
                file_type=self.file_type,
                extraction_succeeded=False,
                extraction_notes=(note,),
            ),
            content="",
            blocks=(),
        )


class ExcelNormalizer:
    """Normalize Excel workbooks while retaining workbook and sheet context."""

    file_type = "xlsx"

    def __init__(self) -> None:
        self._markdown_normalizer = MarkdownNormalizer()

    def normalize_file(self, path: str | Path) -> NormalizedDocument:
        """Create a context-preserving Markdown representation of an Excel workbook."""

        source_path = Path(path)
        if source_path.suffix.lower() != ".xlsx":
            raise ValueError("ExcelNormalizer only accepts .xlsx files")

        try:
            workbook = load_workbook(source_path, data_only=False)
        except Exception as error:
            return self._failed_document(source_path, f"Excel extraction failed: {error}")

        formula_count = 0
        merged_range_count = 0
        image_count = 0
        markdown_lines = [f"# Workbook: {source_path.stem}", ""]
        table_locations: list[tuple[str, int, int, str]] = []

        for worksheet in workbook.worksheets:
            formula_count += sum(
                1
                for row in worksheet.iter_rows()
                for cell in row
                if cell.data_type == "f"
            )
            merged_range_count += len(worksheet.merged_cells.ranges)
            image_count += len(worksheet._images)
            worksheet_lines, worksheet_table_locations = self._normalize_worksheet(worksheet)
            markdown_lines.extend(worksheet_lines)
            table_locations.extend(worksheet_table_locations)

        content = "\n".join(markdown_lines).rstrip() + "\n"
        normalized_markdown = self._markdown_normalizer.normalize_text(
            source_path.name,
            content,
            source_label="Normalized Excel content",
        )
        return NormalizedDocument(
            document=DocumentSummary(
                filename=source_path.name,
                file_type=self.file_type,
                extraction_succeeded=True,
                extraction_notes=("Extracted from workbook structure.",),
                source_metadata=(
                    ("sheet_count", str(len(workbook.worksheets))),
                    ("formula_count", str(formula_count)),
                    ("merged_range_count", str(merged_range_count)),
                    ("image_count", str(image_count)),
                ),
            ),
            content=content,
            blocks=self._attach_sheet_context(normalized_markdown.blocks, table_locations),
        )

    def _normalize_worksheet(
        self, worksheet: object
    ) -> tuple[list[str], list[tuple[str, int, int, str]]]:
        lines = [f"## Sheet: {worksheet.title}", ""]
        table_locations: list[tuple[str, int, int, str]] = []
        row_groups = self._non_empty_row_groups(worksheet)
        pending_title: str | None = None

        if not row_groups:
            return [*lines, "_No populated cells found._", ""], table_locations

        for group in row_groups:
            if self._is_title_row(group):
                pending_title = self._row_values(group[0])[0]
                continue

            if pending_title:
                lines.extend([f"### Table: {pending_title}", ""])
                pending_title = None

            lines.extend(self._to_markdown_table(group))
            lines.append("")
            table_locations.append(self._table_location(worksheet.title, group))

        if pending_title:
            lines.extend([f"### Note: {pending_title}", ""])
        return lines, table_locations

    @staticmethod
    def _non_empty_row_groups(worksheet: object) -> list[list[tuple[object, ...]]]:
        groups: list[list[tuple[object, ...]]] = []
        current_group: list[tuple[object, ...]] = []

        for row in worksheet.iter_rows():
            if any(cell.value not in (None, "") for cell in row):
                current_group.append(row)
            elif current_group:
                groups.append(current_group)
                current_group = []

        if current_group:
            groups.append(current_group)
        return groups

    @classmethod
    def _is_title_row(cls, rows: list[tuple[object, ...]]) -> bool:
        return len(rows) == 1 and len(cls._row_values(rows[0])) == 1

    @staticmethod
    def _row_values(row: tuple[object, ...]) -> list[str]:
        values = ["" if cell.value is None else str(cell.value) for cell in row]
        while values and not values[-1]:
            values.pop()
        return values or [""]

    @classmethod
    def _to_markdown_table(cls, rows: list[tuple[object, ...]]) -> list[str]:
        values_by_row = [cls._row_values(row) for row in rows]
        column_count = max(len(values) for values in values_by_row)
        padded_rows = [values + [""] * (column_count - len(values)) for values in values_by_row]
        header, *body = padded_rows

        lines = [cls._markdown_row(header), cls._markdown_row(["---"] * column_count)]
        lines.extend(cls._markdown_row(row) for row in body)
        return lines

    @staticmethod
    def _markdown_row(values: list[str]) -> str:
        escaped = [value.replace("|", "\\|").replace("\n", "<br>") for value in values]
        return f"| {' | '.join(escaped)} |"

    @staticmethod
    def _table_location(
        sheet_name: str, rows: list[tuple[object, ...]]
    ) -> tuple[str, int, int, str]:
        populated_cells = [cell for row in rows for cell in row if cell.value not in (None, "")]
        row_start = min(cell.row for cell in populated_cells)
        row_end = max(cell.row for cell in populated_cells)
        column_start = min(cell.column for cell in populated_cells)
        column_end = max(cell.column for cell in populated_cells)
        cell_range = (
            f"{get_column_letter(column_start)}{row_start}:"
            f"{get_column_letter(column_end)}{row_end}"
        )
        return sheet_name, row_start, row_end, cell_range

    @staticmethod
    def _attach_sheet_context(
        blocks: tuple[Block, ...], table_locations: list[tuple[str, int, int, str]]
    ) -> tuple[Block, ...]:
        current_sheet: str | None = None
        table_index = 0
        contextualized_blocks: list[Block] = []

        for block in blocks:
            if block.type is BlockType.HEADING and block.heading_level == 2:
                sheet_prefix = "Sheet: "
                if block.text.startswith(sheet_prefix):
                    current_sheet = block.text.removeprefix(sheet_prefix)

            location = replace(block.location, sheet_name=current_sheet)
            if block.type is BlockType.TABLE and table_index < len(table_locations):
                sheet_name, _, _, cell_range = table_locations[table_index]
                location = replace(
                    location,
                    sheet_name=sheet_name,
                    line_start=None,
                    line_end=None,
                    cell_range=cell_range,
                )
                table_index += 1
            contextualized_blocks.append(replace(block, location=location))

        return tuple(contextualized_blocks)

    def _failed_document(self, source_path: Path, note: str) -> NormalizedDocument:
        return NormalizedDocument(
            document=DocumentSummary(
                filename=source_path.name,
                file_type=self.file_type,
                extraction_succeeded=False,
                extraction_notes=(note,),
            ),
            content="",
            blocks=(),
        )
