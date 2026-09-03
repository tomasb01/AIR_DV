"""Tests for deterministic Excel context and row-density findings."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from air_dv.checks.excel import ExcelCheck, ExcelCheckConfig
from air_dv.models import Block, BlockType, DocumentSummary, NormalizedDocument, SourceLocation
from air_dv.normalization import ExcelNormalizer
from tests.fixture_builders import create_context_workbook


class ExcelCheckTests(unittest.TestCase):
    def setUp(self) -> None:
        self.check = ExcelCheck()

    def test_accepts_contextualized_and_titled_tables(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            source_path = Path(temporary_directory, "context.xlsx")
            create_context_workbook(source_path)
            document = ExcelNormalizer().normalize_file(source_path)

        self.assertEqual(self.check.run(document), [])

    def test_reports_missing_sheet_context_as_a_platform_limitation(self) -> None:
        document = NormalizedDocument(
            document=DocumentSummary("lost-context.xlsx", "xlsx", True),
            content="| ID | Name |\n| --- | --- |\n| 1 | Rose |\n",
            blocks=(
                Block(
                    type=BlockType.TABLE,
                    text="| ID | Name |\n| --- | --- |\n| 1 | Rose |",
                    location=SourceLocation(label="Normalized Excel content", line_start=1),
                ),
            ),
        )

        findings = self.check.run(document)

        self.assertEqual(
            [finding.id for finding in findings],
            ["excel-table-missing-sheet-context", "excel-table-has-no-detectable-purpose"],
        )
        self.assertEqual(findings[0].owner.value, "platform_team")
        self.assertEqual(findings[0].category.value, "ingestion_limitation")

    def test_reports_an_untitled_table_as_a_shared_issue(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            source_path = Path(temporary_directory, "untitled.xlsx")
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Data"
            worksheet.append(["ID", "Name"])
            worksheet.append([1, "Rose"])
            workbook.save(source_path)
            document = ExcelNormalizer().normalize_file(source_path)

        findings = self.check.run(document)

        self.assertEqual([finding.id for finding in findings], ["excel-table-has-no-detectable-purpose"])
        self.assertEqual(findings[0].owner.value, "shared")

    def test_reports_only_rows_over_the_configured_density_limit(self) -> None:
        row_at_limit = "x" * 40
        row_over_limit = "y" * 41
        document = NormalizedDocument(
            document=DocumentSummary("dense.xlsx", "xlsx", True),
            content="",
            blocks=(
                Block(
                    type=BlockType.HEADING,
                    text="Table: Records",
                    heading_level=3,
                    location=SourceLocation(label="Normalized Excel content", line_start=1, sheet_name="Data"),
                ),
                Block(
                    type=BlockType.TABLE,
                    text=f"{row_at_limit}\n{row_over_limit}",
                    location=SourceLocation(label="Normalized Excel content", line_start=3, sheet_name="Data"),
                ),
            ),
        )

        findings = ExcelCheck(ExcelCheckConfig(maximum_characters_per_table_row=40)).run(document)

        self.assertEqual([finding.id for finding in findings], ["excel-row-exceeds-recommended-density"])
        self.assertIn("41 characters", findings[0].why_it_matters)

    def test_ignores_non_excel_documents(self) -> None:
        document = NormalizedDocument(
            document=DocumentSummary("notes.md", "md", True),
            content="| ID |\n| --- |\n| 1 |\n",
            blocks=(),
        )

        self.assertEqual(self.check.run(document), [])
