"""Tests for context-preserving Excel normalization."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from air_dv.models import BlockType
from air_dv.normalization import ExcelNormalizer


class ExcelNormalizerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.normalizer = ExcelNormalizer()

    def test_preserves_workbook_sheet_and_table_context(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            source_path = Path(temporary_directory, "oil-catalog.xlsx")
            workbook = Workbook()
            catalog = workbook.active
            catalog.title = "Catalog"
            catalog.append(["Essential oils"])
            catalog.append([])
            catalog.append(["ID", "Name", "Effect"])
            catalog.append([1, "Rose", "Calming"])
            catalog.merge_cells("A1:C1")

            tags = workbook.create_sheet("Tags")
            tags.append(["Tag assignments"])
            tags.append([])
            tags.append(["Tag", "Count", "Calculated total"])
            tags.append(["Calming", 1, "=B4+1"])
            workbook.save(source_path)

            document = self.normalizer.normalize_file(source_path)

        self.assertTrue(document.document.extraction_succeeded)
        self.assertEqual(document.document.file_type, "xlsx")
        self.assertIn("# Workbook: oil-catalog", document.content)
        self.assertIn("## Sheet: Catalog", document.content)
        self.assertIn("### Table: Essential oils", document.content)
        self.assertIn("## Sheet: Tags", document.content)
        self.assertIn("### Table: Tag assignments", document.content)
        self.assertIn("| ID | Name | Effect |", document.content)
        self.assertIn("| Tag | Count | Calculated total |", document.content)
        self.assertIn("| Calming | 1 | =B4+1 |", document.content)
        self.assertEqual(
            dict(document.document.source_metadata),
            {
                "sheet_count": "2",
                "formula_count": "1",
                "merged_range_count": "1",
                "image_count": "0",
            },
        )
        tables = [block for block in document.blocks if block.type is BlockType.TABLE]
        self.assertEqual([table.location.sheet_name for table in tables], ["Catalog", "Tags"])
        self.assertEqual([table.location.cell_range for table in tables], ["A3:C4", "A3:C4"])

    def test_marks_empty_sheets_in_normalized_content(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            source_path = Path(temporary_directory, "empty.xlsx")
            workbook = Workbook()
            workbook.active.title = "Empty"
            workbook.save(source_path)

            document = self.normalizer.normalize_file(source_path)

        self.assertIn("## Sheet: Empty", document.content)
        self.assertIn("_No populated cells found._", document.content)

    def test_rejects_non_excel_files(self) -> None:
        with self.assertRaisesRegex(ValueError, "only accepts .xlsx"):
            self.normalizer.normalize_file("notes.md")
