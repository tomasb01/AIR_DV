"""Reproducible builders for binary test fixtures without business data."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook


def create_visual_title_only_docx(path: Path) -> None:
    """Create a minimal valid DOCX whose bold visual title has no heading style."""

    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""
    relationships = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""
    document = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Visual title only</w:t></w:r></w:p>
    <w:p><w:r><w:t>Body content without a semantic heading style.</w:t></w:r></w:p>
    <w:sectPr/>
  </w:body>
</w:document>"""

    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document)


def create_context_workbook(path: Path) -> None:
    """Create a workbook with metadata that a normalizer must preserve."""

    workbook = Workbook()
    catalog = workbook.active
    catalog.title = "Catalog"
    catalog.append(["Essential oils"])
    catalog.append([])
    catalog.append(["ID", "Name", "Effect"])
    catalog.append([1, "Rose", "Calming"])
    catalog.merge_cells("A1:C1")

    tags = workbook.create_sheet("Tags")
    tags.append(["Tag", "Count"])
    tags.append(["Calming", 1])
    tags["C1"] = "=B2+1"
    workbook.save(path)
